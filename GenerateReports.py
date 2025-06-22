import openpyxl
import pyodbc
import pandas as pd
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
import json
import warnings

warnings.filterwarnings("ignore")

with open('config.json', 'r') as config_file:
    config = json.load(config_file)

database_config = config['staging_db']
doc_config = config['staging_directory']

# Database
_server = database_config['server']
_database = database_config['database']
_user = database_config['user']
_pw = database_config['password']
conn_str = f'DRIVER={{ODBC Driver 17 for SQL Server}};\
            SERVER={_server};\
            DATABASE={_database};\
            UID={_user};\
            PWD={_pw};\
            CHARSET=utf8;COLLATION=Arabic_CI_AI'

def _query_data(_query):
    _conn = pyodbc.connect(conn_str)
    cursor = _conn.cursor()
    cursor.execute(_query)
    results = cursor.fetchall()
    cursor.close()
    data = pd.read_sql_query(_query, _conn)
    df = pd.DataFrame(data)
    _conn.close()
    return df

def create_excel_file(path, sheet_data):

    workbook = Workbook()
    for sheet_name, df in sheet_data.items():
        if sheet_name in workbook.sheetnames:
            workbook.remove(workbook[sheet_name])

        sheet = workbook.create_sheet(title=sheet_name)

        for row in dataframe_to_rows(df, index=False, header=True):
            sheet.append(row)

    default_sheet = workbook.get_sheet_by_name('Sheet')
    workbook.remove(default_sheet)
    workbook.save(path)

partition = doc_config['partition']
folder = doc_config['Folder']
path = f"D:\\MOF_XMigration\\Migration_Reports.xlsx"


successed = _query_data("SELECT DISTINCT T2._Source, T1._Entity, T1._Type \
                        FROM Trace_SPID T1 \
                        LEFT JOIN Mapping.Unified_Mapping T2 \
                        ON T1._Entity = T2.Destination AND T1._Type = T2.Table_Type \
                        ORDER BY _Entity, _Type")

failed = _query_data("SELECT DISTINCT T2._Source, T1._List, T1._Type, T1._Error \
                        FROM Trace_Errors T1 \
                        INNER JOIN Mapping.Unified_Mapping T2 \
                        ON T1._List = T2.Destination")

notfounded_doc = _query_data("SELECT DISTINCT T1.*, T2.ProjectName, T2.EnterpriseProjectTypeName \
                                FROM Not_Founded_DocPSList T1 \
                                INNER JOIN Pre.ProjectsByTaples T2 \
                                ON T1.PID = T2.ProjectId \
                                ORDER BY T1.PID, T1._Name")

notfounded_doc = _query_data("SELECT DISTINCT T1.*, T2.ProjectName, T2.EnterpriseProjectTypeName \
                                FROM Not_Founded_DocPSList T1 \
                                INNER JOIN Pre.ProjectsByTaples T2 \
                                ON T1.PID = T2.ProjectId \
                                ORDER BY T1.PID, T1._Name")

notfounded_lists = _query_data("SELECT DISTINCT T1.*, T2.ProjectName, T2.EnterpriseProjectTypeName \
                                FROM Not_Founded_PSList T1 \
                                INNER JOIN Pre.ProjectsByTaples T2 \
                                ON T1.PID = T2.ProjectId \
                                ORDER BY T1.PID, T1._Name")

notfounded_rootlists = _query_data("SELECT * \
                                    FROM Not_Founded_Lists")

sheet_data = {'Successfully Migrated': successed, 
              'Failed': failed, 
              'NotFoundedDocuments_ProjectSites': notfounded_doc,
              'NotFoundedLists_ProjectSites': notfounded_lists,
              'NotFoundedLists_RootLvl': notfounded_rootlists}

create_excel_file(path, sheet_data)